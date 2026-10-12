import os
import sys
import re
import time
import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Ensure ffmpeg in PATH
os.environ['PATH'] = '/opt/homebrew/bin:' + os.environ.get('PATH', '')

# Load environment variables from .env
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    env_path = os.path.join(os.path.dirname(__file__), '.env')
    if os.path.exists(env_path):
        with open(env_path, 'r', encoding='utf-8') as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith('#') and '=' in line:
                    k, v = line.split('=', 1)
                    os.environ.setdefault(k.strip(), v.strip())

from gemini_transcriber import get_gemini_client, transcribe_audio_gemini
from transliterator import hindi_to_roman


def natural_sort_key(filename):
    """Sort filenames naturally: 1.wav, 1_1.wav, 2.wav, 2_1.wav, 10.wav, etc."""
    base = filename.replace('.wav', '')
    parts = re.split(r'[_.]', base)
    key = []
    for p in parts:
        if p.isdigit():
            key.append(int(p))
        else:
            key.append(p)
    return key

def load_template_context(template_path):
    """Extract stimuli context hints from CBCS template file."""
    if not os.path.exists(template_path):
        return {}
    
    mapping = {}
    try:
        wb = openpyxl.load_workbook(template_path, data_only=True)
        ws = wb['Sheet1']
        for r in range(2, ws.max_row + 1):
            lbl = ws.cell(r, 1).value
            sent = ws.cell(r, 7).value
            if lbl is not None and sent is not None:
                lbl_str = str(lbl).strip()
                # Clean (with pause), (little fumble), etc. from template reference so hint provides semantic context without biasing pause detection
                clean_sent = re.sub(r'\(.*?\)', '', str(sent)).strip()
                clean_sent = re.sub(r'\s+', ' ', clean_sent)
                mapping[lbl_str] = clean_sent
    except Exception as e:
        print(f"Notice: Could not load template context: {e}")
    return mapping

def process_single_audio(client, audio_dir, filename, template_mapping):
    """Processes a single audio file with Gemini."""
    file_path = os.path.join(audio_dir, filename)
    label = filename.replace('.wav', '')
    template_hint = template_mapping.get(str(label), "")
    
    t0 = time.time()
    res = transcribe_audio_gemini(client, file_path, label, template_hint)
    elapsed = time.time() - t0
    
    if res:
        onset = res.get('onset', 0.0)
        offset = res.get('offset', 0.0)
        onset_timing_str = f"{onset:.2f} - {offset:.2f}"
        roman = res.get('roman', '')
        hindi = res.get('hindi', '')
        pause_detected = res.get('pause_detected', False)
        sentence_structure = res.get('sentence_structure', 'NA')
        patient_position = res.get('patient_position', 'NA')
    else:
        onset_timing_str = "Error"
        roman = ""
        hindi = ""
        pause_detected = False
        sentence_structure = "NA"
        patient_position = "NA"
        
    val_label = label  # keep original string label (e.g. '1_1')
    return {
        'label': val_label,
        'filename': filename,
        'onset_timing': onset_timing_str,
        'roman': roman,
        'sentence_structure': sentence_structure,
        'patient_position': patient_position,
        'hindi': hindi,
        'pause': pause_detected,
        'elapsed': elapsed
    }

def transcribe_directory(audio_dir, output_excel, template_file, gemini_api_key=None, max_workers=4):
    api_key = gemini_api_key or os.environ.get('GEMINI_API_KEY')
    if not api_key:
        print("Error: GEMINI_API_KEY not found! Please set GEMINI_API_KEY in .env or pass --gemini-api-key")
        return
    client = get_gemini_client(api_key)
    
    print(f"\n=======================================================")
    print(f" CBCS Multimodal AI Transcription & Timing Pipeline")
    print(f"=======================================================")
    print(f"Audio Source : {audio_dir}")
    print(f"Output Excel : {output_excel}")
    print(f"Template Ref : {template_file}")
    print(f"AI Engine    : Gemini Flash Lite (Low Cost, High Speed)")
    print(f"Onset Method : Triple-Layer Hybrid (VAD + RMS + Gemini)")
    print(f"Concurrency  : {max_workers} worker threads")
    print(f"=======================================================\n")
    
    if not os.path.isdir(audio_dir):
        print(f"Error: Directory {audio_dir} does not exist!")
        return

    wav_files = [f for f in os.listdir(audio_dir) if f.endswith('.wav')]
    wav_files.sort(key=natural_sort_key)
    total = len(wav_files)
    print(f"Found {total} .wav audio files.")
    if total == 0:
        return

    template_mapping = load_template_context(template_file)
    print(f"Loaded {len(template_mapping)} stimuli reference context items.\n")

    # Prepare Workbook
    headers = [
        'RECORDING_SESSION_LABEL',
        'INDEX',
        'Condition I',
        'Condition II',
        'Story',
        'Onset timing',
        'Sentence ',
        'Sentence Structure',
        'Patient Position',
        'Sentence (Hindi)'
    ]
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(headers)

    header_font = Font(name='Calibri', size=11, bold=True, color='000000')
    header_fill = PatternFill(start_color='E8EEF5', end_color='E8EEF5', fill_type='solid')
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Execute concurrent transcription
    print(f"Transcribing {total} audio recordings concurrently with Gemini...")
    start_batch = time.time()
    results_map = {}
    
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        future_to_file = {
            executor.submit(process_single_audio, client, audio_dir, f, template_mapping): f 
            for f in wav_files
        }
        
        completed_count = 0
        for future in as_completed(future_to_file):
            completed_count += 1
            f = future_to_file[future]
            try:
                data = future.result()
                results_map[f] = data
                pause_tag = " [PAUSE]" if data['pause'] else ""
                print(f"[{completed_count:2d}/{total}] {data['filename']:<10} | {data['onset_timing']:<15} | {data['sentence_structure']:<4} | {data['patient_position']:<5} | {data['elapsed']:.2f}s | {data['roman']}{pause_tag}")
            except Exception as e:
                print(f"Error on {f}: {e}")

    # Write in natural sorted order
    for f in wav_files:
        if f in results_map:
            data = results_map[f]
            row_data = [
                data['label'],              # Col A: RECORDING_SESSION_LABEL
                None,                       # Col B: INDEX
                None,                       # Col C: Condition I
                None,                       # Col D: Condition II
                None,                       # Col E: Story
                data['onset_timing'],       # Col F: Onset timing
                data['roman'],              # Col G: Sentence 
                data['sentence_structure'], # Col H: Sentence Structure
                data['patient_position'],   # Col I: Patient Position
                data['hindi']               # Col J: Sentence (Hindi)
            ]
            ws.append(row_data)

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(output_excel)
    total_time = time.time() - start_batch
    print(f"\n=======================================================")
    print(f" All {total} Audios Transcribed in {total_time:.2f} seconds!")
    print(f" Output saved to: {output_excel}")
    print(f"=======================================================\n")

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Transcribe CBCS audio recordings using Gemini API.")
    parser.add_argument('--audio-dir', default='/Users/dhananjaykdhani/Downloads/Data_s4_hindi_p3/recorded_audio_Chandan')
    parser.add_argument('--output-excel', default='/Users/dhananjaykdhani/Desktop/Transcript/CBCS_Audio File Chandan.xlsx')
    parser.add_argument('--template', default='/Users/dhananjaykdhani/Desktop/Transcript/CBCS_Audio File Anshu.xlsx')
    parser.add_argument('--gemini-api-key', default=None, help='Gemini API key (defaults to GEMINI_API_KEY from .env)')
    parser.add_argument('--workers', type=int, default=4)
    args = parser.parse_args()
    
    transcribe_directory(args.audio_dir, args.output_excel, args.template, args.gemini_api_key, args.workers)
